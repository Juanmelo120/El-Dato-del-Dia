import glob, json, os, re, docx
R = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Títulos y temas originales (para los "=")
src = open(os.path.join(R, 'contenido/carruseles_original.txt'), encoding='utf-8').read().split('\n')
orig = {}
for i, l in enumerate(src):
    m = re.match(r'CARRUSEL (\d+) — (.*)', l)
    if m:
        n = int(m.group(1)); t = src[i+1].replace('Título: ', '').strip()
        t = re.sub(r'\s+I$', '', t)
        orig[n] = (m.group(2).strip(), t)

car = {}
for f in sorted(glob.glob(os.path.join(R, 'contenido/carruseles/*.txt'))):
    cur = None
    for l in open(f, encoding='utf-8'):
        l = l.rstrip('\n')
        if not l.strip(): continue
        if l.startswith('#'):
            p = l[1:].split('|'); n = int(p[0])
            tema = orig[n][0] if p[1] == '=' else p[1]
            tit = orig[n][1] if p[2] == '=' else p[2]
            car[n] = dict(n=n, tema=tema, titulo=tit, gancho=p[3].strip(), datos=[]); cur = n; continue
        t, g = [x.strip() for x in l.split('|', 1)]
        car[cur]['datos'].append(dict(titular=t, gancho=g))
assert sorted(car) == list(range(1, 501)) and all(len(c['datos']) == 10 for c in car.values())
for c in car.values():  # los datos están escritos del 10 al 1
    for i, d in enumerate(c['datos']): d['n'] = 10 - i

json.dump([car[n] for n in sorted(car)], open(os.path.join(R, 'src/data/carruseles.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

D = docx.Document()
st = D.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
H = lambda t, l=1: D.add_heading(t, level=l)
def P(t, bold=False):
    p = D.add_paragraph(); r = p.add_run(t); r.bold = bold; return p
def B(t): D.add_paragraph(t, style='List Bullet')

H('500 CARRUSELES — "EL DATO DEL DÍA"', 0)
P('Versión 3 · 500 carruseles únicos · textos fijos iguales a la web', True)
H('Qué cambió en esta versión')
B('Los 500 carruseles son distintos: se eliminaron las repeticiones (antes 001–020 se repetían en 021–040, 041–060, 061–080 y 081–100). Son 5.000 datos y ninguno se repite.')
B('Cada diapositiva trae su texto exacto (TITULAR + GANCHO). Gemini ya no inventa el dato: solo crea la imagen y copia el texto. Así lo que sale en TikTok coincide con lo que se lee en la web.')
B('El GANCHO es una frase intrigante que no lo cuenta todo, para que la gente quiera ir a la página a leer el porqué.')
B('El PIE de cada imagen y el CTA final mandan a la web: "Explicación completa en la web · link en el perfil · Dato #NNN".')
B('En la web, la página /tiktok busca por número: si alguien escribe 127, llega directo al artículo del carrusel 127.')
H('Reglas maestras para Gemini (aplican a todas las imágenes)')
for r in ['Formato vertical 9:16 (1080x1920), optimizado para TikTok.',
 'Copia los TEXTOS EXACTOS tal cual, entre comillas: sin cambiar palabras, sin resumir y sin añadir datos. Revisa la ortografía y las tildes.',
 'La imagen debe mostrar exactamente lo que dice el TITULAR, no una imagen genérica del tema.',
 'No repitas sujeto principal, escenario, encuadre ni composición entre las diapositivas del mismo carrusel.',
 'Estilo hiperrealista y cinematográfico, con la paleta de la marca: fondo nocturno morado, acentos dorado (#FFC83D), cian (#3FD9FF) y magenta neón (#E24BFF).',
 'TITULAR grande en blanco o dorado; GANCHO más pequeño en cian; PIE pequeño abajo, legible. El texto no debe tapar el elemento principal.',
 'Marca discreta "EL DATO DEL DÍA" arriba.',
 'Sin interfaz de TikTok, teléfono, botones, likes, comentarios ni capturas de pantalla.']:
    B(r)
H('Estructura de cada carrusel')
P('Portada → Dato 10 → 9 → 8 → 7 → 6 → 5 → 4 → 3 → 2 → Dato 1 → CTA final (12 imágenes).')
P('Aviso de calidad: los datos se escribieron a partir de conocimiento general y se revisaron para evitar mitos y duplicados. Antes de publicar cada carrusel, su artículo web se redacta con fuentes y, si algún dato necesita matiz o corrección, se corrige también aquí.')

EST = 'Estilo hiperrealista y cinematográfico con la paleta de la marca (morado nocturno, dorado, cian y magenta neón). Texto grande, legible y sin errores. Marca discreta "EL DATO DEL DÍA". Sin interfaz de TikTok, teléfono, botones ni likes.'
for n in sorted(car):
    c = car[n]; num = f'{n:03d}'
    pie = f'Explicación completa en la web · link en el perfil · Dato #{num}'
    D.add_page_break()
    H(f'CARRUSEL {num} — {c["tema"]}', 1)
    P(f'Título: {c["titulo"]}', True)
    P(f'Gancho de portada: {c["gancho"]}')
    H('PROMPT 1 — PORTADA', 2)
    lista = '; '.join(d['titular'] for d in c['datos'])
    P(f'Crea una imagen vertical 9:16 (1080x1920) para "El Dato del Día". PORTADA. Tema: {c["tema"]}. ESCENA: portada cinematográfica que integre, de forma equilibrada y sin un protagonista único, elementos visuales de estos 10 datos: {lista}. Deja espacio limpio para el texto. '
      f'TEXTOS EXACTOS: TITULAR: "{c["titulo"].upper()}". GANCHO: "{c["gancho"]}". PIE: "{pie}". {EST}')
    for i, d in enumerate(c['datos'], 2):
        H(f'PROMPT {i} — DATO {d["n"]}', 2)
        P(f'Crea una imagen vertical 9:16 (1080x1920) para "El Dato del Día". DATO {d["n"]} de 10. Tema: {c["tema"]}. ESCENA: una escena específica que muestre literalmente este dato: "{d["titular"]}". '
          f'TEXTOS EXACTOS: TITULAR: "{d["titular"].upper()}". GANCHO: "{d["gancho"]}". PIE: "{pie}". {EST}')
    H('PROMPT 12 — CTA FINAL (llevar a la web)', 2)
    P(f'Crea una imagen vertical 9:16 (1080x1920), última diapositiva de "El Dato del Día". Tema: {c["tema"]}. ESCENA: original y distinta a las anteriores, relacionada con el tema y con la idea de descubrir; incluye una bombilla brillante como símbolo de la marca. '
      f'TEXTOS EXACTOS: TITULAR: "¿QUIERES SABER EL PORQUÉ?". TEXTO: "Cada uno de estos 10 datos tiene su explicación completa, con fuentes, en nuestra web". DESTACADO: "LINK EN EL PERFIL → DATO #{num}". PIE: "❤️ Dale like · 💬 Comenta el número que más te sorprendió · ↗️ Compártelo". {EST}')
    P(f'Descripción para TikTok: {c["gancho"]} 🤯 {c["titulo"]}. ¿Cuál te sorprendió más? La explicación completa de los 10 datos, con fuentes, está en nuestra web: link en el perfil → busca el Dato #{num}. #fyp #parati #datoscuriosos #sabiasque #curiosidades #elDatodelDia')

D.save(os.path.join(R, '500_Carruseles_El_Dato_del_Dia_v3.docx'))
print('ok', len(car))
